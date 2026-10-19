# -*- coding: utf-8 -*-
"""get_training_plan_doc（培养方案附件正文，F04 扩展 2026-10-18）合成回归。

全部使用假传输/假会话/临时目录，不触网、无真实凭据。
覆盖：成功下载落盘、动态 id 取自本会话列表、方案名不匹配、
列表/下载会话失效、下载业务拒绝（出错页）、失败保旧（partial stale_file）。
"""
from __future__ import annotations

import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.adapters.campus.jwweb import fetch_training_plan_doc
from src.auth import session as S
from src.service.jwweb_service import JWWebService
from src.storage.store import SnapshotStore

_LIST_ITEM = {"pyfamc": "2024示例专业", "bbh": "2023",
              "pyfaxx_id": "ABCDEF0123456789ABCDEF0123456789-2025",
              "zdxf": "160.0", "hasPyfaxxFj": "1",
              "njdm_id": "2025", "zymc": "示例专业"}


def _docx_bytes() -> bytes:
    """合成合法 DOCX（发布前校验要求 ZIP/DOCX 成员齐全）。"""
    doc = Document()
    doc.add_paragraph("培养方案正文")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "课程"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paged_body(items):
    return json.dumps({"currentPage": 1, "items": items, "totalCount": len(items),
                       "totalPage": 1, "totalResult": len(items)}, ensure_ascii=False)


class _FakeTransport:
    """脚本化假传输：按请求顺序弹出；记录调用供断言。"""

    def __init__(self, script):
        self.script = list(script)
        self.calls = []

    def get(self, url, headers=None, timeout=20):
        self.calls.append(("GET", url, None))
        return self.script.pop(0)

    def get_binary(self, url, headers=None, timeout=30):
        self.calls.append(("GETBIN", url, None))
        return self.script.pop(0)

    def post_form(self, url, form=None, headers=None, timeout=20):
        self.calls.append(("POST", url, dict(form or {})))
        return self.script.pop(0)

    def export_cookies(self):
        return []

    def set_cookies(self, c):
        pass


def _r_text(text, status=200):
    return {"kind": "response", "status": status, "text": text}


def _r_bin(body, status=200):
    return {"kind": "response", "status": status, "body": body}


class TrainingPlanDocServiceTest(unittest.TestCase):
    def _mk_service(self, tmp, fake=None):
        S.save_session(tmp, "acct-tpd", {"zjgsu_token": "t", "user_type": "2",
                                         "jwglxt_cookies": [{"name": "JSESSIONID", "value": "x",
                                                             "domain": "jwxt.zjgsu.edu.cn",
                                                             "path": "/", "secure": False,
                                                             "expires": None, "discard": True}]})
        S.set_active(tmp, "acct-tpd")
        return JWWebService(SnapshotStore(tmp / "c.sqlite3"), tmp, jwweb_transport=fake)

    def _ok_script(self, body: bytes | None = None):
        return [
            _r_text("<html>index</html>"),          # 列表激活
            _r_text(_paged_body([_LIST_ITEM])),     # 列表
            _r_text("<html>index</html>"),          # 下载前激活
            _r_bin(body if body is not None else _docx_bytes()),  # 下载
        ]

    def test_ok_download_writes_docx_and_entry(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_ok_"))
        body = _docx_bytes()
        ft = _FakeTransport(self._ok_script(body))
        svc = self._mk_service(tmp, ft)
        out = svc.get_training_plan_doc("2024示例专业", "2025")
        self.assertEqual(out["status"], "ok", json.dumps(out.get("issues", []), ensure_ascii=False))
        self.assertEqual(len(out["items"]), 1)
        e = out["items"][0]
        self.assertEqual(e["kind"], "training_plan_doc")
        self.assertEqual(e["plan_name"], "2024示例专业")
        self.assertEqual(e["size_bytes"], len(body))
        docx = Path(e["file_path"])
        self.assertTrue(docx.exists())
        self.assertEqual(docx.read_bytes()[:2], b"PK")
        self.assertIn("training_plan_docs", e["file_path"])
        self.assertIn("acct-tpd", e["file_path"])
        self.assertIn("2025", e["file_path"])
        # 下载 id 必须取自本会话列表返回的 pyfaxx_id（动态 id 规则）
        dl = [c for c in ft.calls if c[0] == "GETBIN"]
        self.assertEqual(len(dl), 1)
        self.assertIn("id=ABCDEF0123456789ABCDEF0123456789", dl[0][1])

    def test_dynamic_id_from_list_not_stale(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_dyn_"))
        ft = _FakeTransport(self._ok_script())
        svc = self._mk_service(tmp, ft)
        svc.get_training_plan_doc("2024示例专业", "2025")
        dl = [c for c in ft.calls if c[0] == "GETBIN"][0][1]
        # 旧会话 id（如 FEEDC0DE…）不得出现——每次用本会话列表 id
        self.assertNotIn("FEEDC0DE", dl)

    def test_plan_name_not_found_is_unavailable(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_nf_"))
        ft = _FakeTransport([
            _r_text("<html>index</html>"),
            _r_text(_paged_body([{**_LIST_ITEM, "pyfamc": "2024另一示例专业"}])),
        ])
        svc = self._mk_service(tmp, ft)
        out = svc.get_training_plan_doc("2024示例专业", "2025")
        self.assertEqual(out["status"], "unavailable")
        self.assertTrue(any("完全同名" in i["safe_message"]
                            for i in out["issues"]))

    def test_list_session_expired_is_auth_required(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_exp_"))
        ft = _FakeTransport([
            _r_text("<html>index</html>"),
            _r_text("<html>统一身份认证 cas/login 请登录</html>"),
        ])
        svc = self._mk_service(tmp, ft)
        out = svc.get_training_plan_doc("2024示例专业", "2025")
        self.assertEqual(out["status"], "unavailable")
        self.assertTrue(any(i["code"] == "auth_required" for i in out["issues"]))

    def test_doc_endpoint_upstream_error_page(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_err_"))
        ft = _FakeTransport([
            _r_text("<html>index</html>"),
            _r_text(_paged_body([_LIST_ITEM])),
            _r_text("<html>index</html>"),
            _r_text("<html><title>错误提示</title>出错啦</html>"),
        ])
        svc = self._mk_service(tmp, ft)
        out = svc.get_training_plan_doc("2024示例专业", "2025")
        self.assertEqual(out["status"], "unavailable")
        self.assertTrue(any("附件下载失败" in i["safe_message"]
                            for i in out["issues"]))

    def test_failure_keeps_old_file_as_partial(self):
        tmp = Path(tempfile.mkdtemp(prefix="tpd_old_"))
        body = _docx_bytes()
        svc = self._mk_service(tmp, _FakeTransport([
            _r_text("<html>index</html>"),
            _r_text(_paged_body([_LIST_ITEM])),
            _r_text("<html>index</html>"),
            _r_text("<html><title>错误提示</title>出错啦</html>"),
        ]))
        paths = svc._plan_doc_paths("acct-tpd", "2025", "2024示例专业")
        old_dir = paths["dir"]
        old_dir.mkdir(parents=True)
        old_file = paths["docx"]
        old_file.write_bytes(body)
        paths["meta"].write_text(json.dumps({
            "schema": 1, "account": "acct-tpd", "apply_grade": "2025",
            "plan_name": "2024示例专业", "version": "2023",
            "total_credit": "160.0", "size_bytes": len(body),
            "sha256": hashlib.sha256(body).hexdigest(),
            "text_chars": None, "text_complete": False, "text_preview": None,
            "text_path": None, "downloaded_at": "2026-10-18T14:24:00+08:00",
            "source_id": "src.jwglxt",
        }, ensure_ascii=False), encoding="utf-8")
        out = svc.get_training_plan_doc("2024示例专业", "2025")
        self.assertEqual(out["status"], "partial")
        self.assertEqual(len(out["items"]), 1)
        self.assertEqual(out["items"][0]["kind"], "stale_file")
        self.assertEqual(out["items"][0]["size_bytes"], len(body))
        self.assertEqual(old_file.read_bytes(), body)


class TrainingPlanDocAdapterTest(unittest.TestCase):
    """适配器分类：PK 成功 / 登录页 / 出错页 / 非 200 / 网络错误。"""

    def _transport(self, script):
        ft = _FakeTransport(script)
        # fetch_training_plan_doc 只用 get / get_binary
        return ft

    def test_ok(self):
        t = self._transport([
            _r_text("<html>index</html>"),
            _r_bin(b"PK\x03\x04data"),
        ])
        res = fetch_training_plan_doc(t, "ABCDEF0123456789ABCDEF0123456789")
        self.assertEqual(res["outcome"], "ok")
        self.assertEqual(res["size"], 8)
        self.assertEqual(res["content_type"], "application/msword")

    def test_session_expired(self):
        t = self._transport([
            _r_text("<html>index</html>"),
            _r_bin("<html>统一身份认证 cas/login</html>".encode("utf-8")),
        ])
        res = fetch_training_plan_doc(t, "ABCDEF0123456789ABCDEF0123456789")
        self.assertEqual(res["outcome"], "session_expired")

    def test_upstream_error_page(self):
        t = self._transport([
            _r_text("<html>index</html>"),
            _r_bin("<html>出错啦！</html>".encode("utf-8")),
        ])
        res = fetch_training_plan_doc(t, "ABCDEF0123456789ABCDEF0123456789")
        self.assertEqual(res["outcome"], "upstream_error")

    def test_http_404(self):
        t = self._transport([
            _r_text("<html>index</html>"),
            _r_bin(b"<html>Sorry, Page Not Found</html>", status=404),
        ])
        res = fetch_training_plan_doc(t, "ABCDEF0123456789ABCDEF0123456789")
        self.assertEqual(res["outcome"], "http_error")

    def test_network_error(self):
        t = self._transport([
            _r_text("<html>index</html>"),
            {"kind": "network_error", "err": "conn refused"},
        ])
        res = fetch_training_plan_doc(t, "ABCDEF0123456789ABCDEF0123456789")
        self.assertEqual(res["outcome"], "network_error")


if __name__ == "__main__":
    unittest.main()
