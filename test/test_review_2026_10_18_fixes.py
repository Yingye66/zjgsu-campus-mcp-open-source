# -*- coding: utf-8 -*-
"""2026-10-18 独立审查 §10.2 修复回归（全部合成，不触网、不读真实 data）。

覆盖：
1. 选座默认拒绝 HTTP 明文发送会话 Cookie（opt-in 才允许）；
2. 选座零容量不除零；
3. 电表非 JSON 受控返回、month 格式校验；
4. 电表园区唯一定位（同名楼栋歧义拒绝 / area_root_id 限定）；
5. 培养方案只接受唯一精确匹配（不替用户选相近方案）；
6. 附件文件键纳入账号+年级，旧件只认元数据+哈希；
7. 附件先校验 ZIP/DOCX 再原子发布，失败保旧；
8. docx 按正文顺序提取（段落/表格交错）；
9. 选座/电表最小观测契约（freshness、partial+as_of、首次失败 unavailable）；
10. 两新工具具备 outputSchema。
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest import mock

from docx import Document

from test.helpers import seed_session, temp_dir

import src.service.electric_service as ES
import src.service.jwweb_service as JWS
import src.service.seat_service as SS
from src.adapters.campus import electric as E
from src.adapters.campus import seat as SEAT
from src.adapters.campus.jwweb import JwwebFetch
from src.auth import session as S
from src.service.jwweb_service import (JWWebService, _extract_docx_text,
                                       _verify_docx_bytes)
from src.service.electric_service import ElectricService
from src.service.seat_service import SeatService
from src.storage.store import SnapshotStore


# ---------------- 合成工具 ----------------

class _FakeResp:
    def __init__(self, status=200, body: bytes = b"", headers=None):
        self.status_code = status
        self.content = body
        self.headers = headers or {}

    @property
    def text(self):
        return self.content.decode("utf-8", "replace")


class _JsonSession:
    """按顺序返回 JSON 的假 requests.Session（只实现 post/get）。"""

    def __init__(self, posts, get_body: bytes = b""):
        self._posts = list(posts)
        self._get_body = get_body
        self.calls = []

    def post(self, url, json=None, data=None, headers=None, timeout=20):
        self.calls.append((url, json if json is not None else data))
        if not self._posts:
            raise AssertionError("假会话响应队列为空：%s" % url)
        return self._posts.pop(0)

    def get(self, url, headers=None, timeout=20, allow_redirects=False):
        self.calls.append((url, None))
        return _FakeResp(200, self._get_body)


def _make_docx(tag: str = "A") -> bytes:
    d = Document()
    d.add_paragraph("标题%s" % tag)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "值"
    d.add_paragraph("标题B")
    t2 = d.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = "条款"
    t2.cell(0, 1).text = "内容"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _seed_jw(data_dir, acct="acct-review"):
    seed_session(data_dir, acct)
    sess = S.load_session(data_dir, acct) or {}
    sess["jwglxt_cookies"] = [{"name": "JSESSIONID", "value": "SYNTH",
                               "domain": "jw.zjgsu.edu.cn", "path": "/"}]
    S.save_session(data_dir, acct, sess)
    return acct


def _plan_rows(*names):
    return [{"pyfamc": n, "pyfaxx_id": "A" * 32, "synj": "2025",
             "bbh": "2023", "zdxf": "168"} for n in names]


def _svc(data_dir) -> JWWebService:
    return JWWebService(SnapshotStore(Path(data_dir) / "cache.sqlite3"), data_dir)


def _message(r: dict) -> str:
    """观测结构里 unavailable 的说明在 issues[].safe_message，partial 另有 reason。"""
    if r.get("reason"):
        return str(r["reason"])
    issues = r.get("issues") or []
    if not issues:
        return ""
    row = issues[0] or {}
    return str(row.get("safe_message") or row.get("message")
               or row.get("detail") or "")


# ---------------- 1. 选座传输安全 ----------------

class SeatTransportTests(unittest.TestCase):
    def test_default_refuses_plaintext_cookie(self):
        with mock.patch.object(SEAT.requests, "get",
                               side_effect=AssertionError("不得发出请求")):
            with mock.patch.dict(os.environ, {SEAT._INSECURE_ENV: ""}):
                res = SEAT.fetch_areas("wechatSESS_ID=SECRET")
        self.assertEqual(res["outcome"], "insecure_transport")
        self.assertNotIn("SECRET", res["diag"])
        self.assertTrue(SEAT.MANUAL_GUIDE)

    def test_opt_in_allows_http_but_keeps_manual_guide(self):
        html = ('<div class="list-group"><a data-url="/index.php/reserve/'
                'layout/libid=6.html?ts=1"><h4>四楼书库 (四楼) '
                '<span class="badge">10/20</span></h4></a></div>').encode()
        with mock.patch.object(SEAT.requests, "get",
                               return_value=_FakeResp(200, html)):
            with mock.patch.dict(os.environ, {SEAT._INSECURE_ENV: "1"}):
                res = SEAT.fetch_areas("wechatSESS_ID=SECRET")
        self.assertEqual(res["outcome"], "ok")
        self.assertEqual(res["areas"][0]["free"], 10)
        self.assertEqual(res["areas"][0]["total"], 20)

    def test_zero_capacity_no_division(self):
        agg = SEAT.aggregate_by_floor([
            {"lib_id": 1, "name": "空区", "floor": "一", "free": 0, "total": 0},
            {"lib_id": 2, "name": "有座", "floor": "二", "free": 3, "total": 4},
        ])
        by_name = {x["name"]: x for x in agg["rank_by_free"]}
        self.assertIsNone(by_name["空区"]["free_rate"])
        self.assertEqual(by_name["有座"]["free_rate"], 0.75)
        # 排序不得因 None 抛异常
        self.assertEqual(len(agg["rank_by_rate"]), 2)


class SeatServiceTests(unittest.TestCase):
    def _prep(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        seed_session(d, "acct-seat")
        sess = S.load_session(d, "acct-seat") or {}
        sess["seat_cookie"] = "wechatSESS_ID=SECRET"
        S.save_session(d, "acct-seat", sess)
        return tmp, d

    def test_insecure_returns_manual_path(self):
        tmp, d = self._prep()
        with tmp:
            with mock.patch.dict(os.environ, {SEAT._INSECURE_ENV: ""}):
                r = SeatService(d).get_seat_summary()
        self.assertEqual(r["status"], "unavailable")
        self.assertIn("insecure_transport", r["reason"])
        self.assertTrue(r["manual_path"])
        self.assertNotIn("SECRET", json.dumps(r, ensure_ascii=False))

    def test_stale_partial_after_success(self):
        tmp, d = self._prep()
        payload = {"areas": [{"lib_id": 6, "name": "四楼书库", "floor": "四",
                              "free": 10, "total": 20}],
                   "by_floor": {"四": {"free": 10, "total": 20, "areas": []}},
                   "rank_by_rate": [{"name": "四楼书库", "free_rate": 0.5}]}
        with tmp:
            svc = SeatService(d)
            with mock.patch.object(SS, "fetch_seat_summary",
                                   return_value={"outcome": "ok", "diag": "",
                                                 "data": payload}):
                r1 = svc.get_seat_summary()
            with mock.patch.object(SS, "fetch_seat_summary",
                                   return_value={"outcome": "network_error",
                                                 "diag": "synthetic", "data": None}):
                r2 = svc.get_seat_summary()
        self.assertEqual(r1["status"], "ok")
        self.assertFalse(r1["freshness"]["stale"])
        self.assertEqual(r2["status"], "partial")
        self.assertTrue(r2["freshness"]["stale"])
        self.assertTrue(r2["as_of"])
        self.assertEqual(r2["areas"], payload["areas"])
        self.assertIn("不代表此刻", r2["reason"])

    def test_first_failure_is_unavailable_not_fake_empty(self):
        tmp, d = self._prep()
        with tmp:
            with mock.patch.object(SS, "fetch_seat_summary",
                                   return_value={"outcome": "parse_error",
                                                 "diag": "synthetic", "data": None}):
                r = SeatService(d).get_seat_summary()
        self.assertEqual(r["status"], "unavailable")
        self.assertNotIn("areas", r)


# ---------------- 2. 电表边界 ----------------

class ElectricAdapterTests(unittest.TestCase):
    def test_bill_non_json_is_controlled(self):
        sess = _JsonSession([_FakeResp(200, b"<html>not json</html>")])
        with mock.patch.object(E, "_bootstrap", return_value=sess):
            res = E.fetch_meter_bill("999999999999")
        self.assertEqual(res["outcome"], "parse_error")
        self.assertIsNone(res["data"])

    def test_invalid_month_rejected_without_network(self):
        with mock.patch.object(E, "_bootstrap",
                               side_effect=AssertionError("不得联网")):
            res = E.fetch_meter_detail("999999999999", month="2026-13")
        self.assertEqual(res["outcome"], "invalid_arguments")
        self.assertIn("YYYY-MM", res["diag"])

    def test_ambiguous_building_across_areas_refused(self):
        sess = _JsonSession([
            _FakeResp(200, json.dumps({"floor": [
                {"BuildingName": "1幢", "BuildingID": "B1"}]}).encode()),
            _FakeResp(200, json.dumps({"floor": [
                {"BuildingName": "1幢", "BuildingID": "B2"}]}).encode()),
        ])
        with mock.patch.object(E, "_ROOT_IDS", ["r1", "r2"]):
            with mock.patch.object(E, "_bootstrap", return_value=sess):
                res = E.resolve_room_meter("1幢", "9楼", "909")
        self.assertEqual(res["outcome"], "ambiguous")
        self.assertIn("area_root_id", res["diag"])
        self.assertEqual(len(sess.calls), 2)

    def test_area_root_id_limits_search_and_resolves(self):
        href_html = b"<script>var x={comAddress: '999999999999'};</script>"
        sess = _JsonSession([
            _FakeResp(200, json.dumps({"floor": [
                {"BuildingName": "1幢", "BuildingID": "B2"}]}).encode()),
            _FakeResp(200, json.dumps({"room": [
                {"BuildingName": "9楼", "BuildingID": "F2"}]}).encode()),
            _FakeResp(200, json.dumps({"room1": [
                {"BuildingName": "909房", "BuildingID": "R2"}]}).encode()),
            _FakeResp(200, json.dumps({"href": "/selectmeter?x=1"}).encode()),
        ], get_body=href_html)
        with mock.patch.object(E, "_ROOT_IDS", ["r1", "r2"]):
            with mock.patch.object(E, "_bootstrap", return_value=sess):
                res = E.resolve_room_meter("1幢", "9楼", "909",
                                           area_root_id="r2")
        self.assertEqual(res["outcome"], "ok")
        self.assertEqual(res["meter"], "999999999999")
        self.assertEqual(res["root_id"], "r2")
        # 只查了 r2 一个园区
        self.assertEqual(sess.calls[0][1], {"bID": "r2"})


class ElectricServiceTests(unittest.TestCase):
    OK = {"outcome": "ok", "diag": "", "data": {
        "meter": "M1", "prepaid_yuan": "35.30", "subsidy_yuan": "0",
        "total_yuan": "35.30", "total_kwh": "10876.33",
        "meter_status": "保电", "reading_time": "2026-10-18 02:44:34",
        "room": "909", "charge_type": "预付费"}}

    def _prep(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        (d / "dorm_meter.json").write_text(json.dumps({
            "meter": "M1", "building": "1幢", "floor": "9楼", "room": "909",
            "area_label": "示例学生公寓", "area_root_id": "001000099",
        }, ensure_ascii=False), encoding="utf-8")
        return tmp, d

    def test_success_has_live_freshness(self):
        tmp, d = self._prep()
        with tmp:
            with mock.patch.object(ES.electric, "fetch_meter_summary",
                                   return_value=self.OK):
                r = ElectricService(d).get_dorm_electric("summary")
        self.assertEqual(r["status"], "ok")
        self.assertEqual(r["freshness"]["source"], "live")
        self.assertFalse(r["freshness"]["stale"])
        self.assertEqual(r["freshness"]["data_age_seconds"], 0)

    def test_failure_returns_stale_partial(self):
        tmp, d = self._prep()
        with tmp:
            svc = ElectricService(d)
            with mock.patch.object(ES.electric, "fetch_meter_summary",
                                   return_value=self.OK):
                svc.get_dorm_electric("summary")
            with mock.patch.object(ES.electric, "fetch_meter_summary",
                                   return_value={"outcome": "network_error",
                                                 "diag": "synthetic", "data": None}):
                r = svc.get_dorm_electric("summary")
        self.assertEqual(r["status"], "partial")
        self.assertTrue(r["freshness"]["stale"])
        self.assertTrue(r["as_of"])
        self.assertEqual(r["data"]["total_yuan"], "35.30")
        self.assertIn("不代表当前", r["reason"])

    def test_first_failure_unavailable(self):
        tmp, d = self._prep()
        with tmp:
            with mock.patch.object(ES.electric, "fetch_meter_summary",
                                   return_value={"outcome": "error",
                                                 "diag": "synthetic", "data": None}):
                r = ElectricService(d).get_dorm_electric("summary")
        self.assertEqual(r["status"], "unavailable")
        self.assertNotIn("data", r)


# ---------------- 3. 培养方案附件 ----------------

class TrainingPlanDocTests(unittest.TestCase):
    def _prep(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        acct = _seed_jw(d)
        return tmp, d, acct

    def _run(self, d, rows, download, plan="2024示例专业", grade="2025"):
        dl = (download if isinstance(download, mock.Mock)
              else mock.Mock(return_value=download))
        with mock.patch.object(JWS, "fetch_training_plans",
                              return_value=JwwebFetch("ok", rows, "", len(rows))):
            with mock.patch.object(JWS, "fetch_training_plan_doc", dl):
                return _svc(d).get_training_plan_doc(plan, grade)

    def test_ambiguous_exact_name_refused(self):
        tmp, d, _ = self._prep()
        rows = _plan_rows("2024示例专业", "2024示例专业")
        no_download = mock.Mock(side_effect=AssertionError("不得下载"))
        with tmp:
            r = self._run(d, rows, no_download)
        no_download.assert_not_called()
        self.assertEqual(r["status"], "unavailable")
        self.assertIn("同名方案", _message(r))

    def test_no_substring_fallback(self):
        tmp, d, _ = self._prep()
        rows = _plan_rows("2024示例专业（专升本）")
        no_download = mock.Mock(side_effect=AssertionError("不得下载"))
        with tmp:
            r = self._run(d, rows, no_download, plan="软件工程")
        no_download.assert_not_called()
        self.assertEqual(r["status"], "unavailable")
        self.assertIn("完全同名", _message(r))
        self.assertIn("相近条目", _message(r))

    def test_grade_mismatch_refused(self):
        tmp, d, _ = self._prep()
        rows = [{"pyfamc": "2024示例专业", "pyfaxx_id": "A" * 32,
                 "synj": "2024", "bbh": "2023", "zdxf": "168"}]
        no_download = mock.Mock(side_effect=AssertionError("不得下载"))
        with tmp:
            r = self._run(d, rows, no_download)
        no_download.assert_not_called()
        self.assertEqual(r["status"], "unavailable")
        self.assertIn("适用年级", _message(r))

    def test_publish_verifies_writes_meta_and_txt(self):
        tmp, d, acct = self._prep()
        body = _make_docx()
        with tmp:
            r = self._run(d, _plan_rows("2024示例专业"),
                          {"outcome": "ok", "body": body, "size": len(body),
                           "content_type": "application/msword"})
            paths = _svc(d)._plan_doc_paths(acct, "2025", "2024示例专业")
            meta = json.loads(paths["meta"].read_text(encoding="utf-8"))
            self.assertTrue(paths["docx"].is_file())
            self.assertTrue(paths["txt"].is_file())
            self.assertEqual(meta["sha256"],
                             hashlib.sha256(body).hexdigest())
        self.assertEqual(r["status"], "ok")
        entry = r["items"][0]
        self.assertEqual(entry["sha256"], hashlib.sha256(body).hexdigest())
        self.assertTrue(entry["text_complete"])
        self.assertEqual(entry["size_bytes"], len(body))

    def test_publish_failure_rolls_back_old_version(self):
        tmp, d, acct = self._prep()
        first = _make_docx("A")
        second = _make_docx("B")
        real_replace = os.replace
        calls = {"n": 0}

        def flaky_replace(src, dst):
            calls["n"] += 1
            if calls["n"] == 2:      # docx 已替换，txt 替换失败 → 必须回滚
                raise OSError("synthetic replace failure")
            return real_replace(src, dst)

        with tmp:
            r1 = self._run(d, _plan_rows("2024示例专业"),
                           {"outcome": "ok", "body": first, "size": len(first),
                            "content_type": "application/msword"})
            self.assertEqual(r1["status"], "ok")
            paths = _svc(d)._plan_doc_paths(acct, "2025", "2024示例专业")
            before_txt = paths["txt"].read_bytes()
            with mock.patch.object(JWS.os, "replace", side_effect=flaky_replace):
                r2 = self._run(d, _plan_rows("2024示例专业"),
                               {"outcome": "ok", "body": second,
                                "size": len(second),
                                "content_type": "application/msword"})
            self.assertEqual(paths["docx"].read_bytes(), first)
            self.assertEqual(paths["txt"].read_bytes(), before_txt)
        self.assertEqual(r2["status"], "partial")

    def test_invalid_pk_never_overwrites_old(self):
        tmp, d, acct = self._prep()
        good = _make_docx()
        with tmp:
            r1 = self._run(d, _plan_rows("2024示例专业"),
                           {"outcome": "ok", "body": good, "size": len(good),
                            "content_type": "application/msword"})
            self.assertEqual(r1["status"], "ok")
            paths = _svc(d)._plan_doc_paths(acct, "2025", "2024示例专业")
            before = paths["docx"].read_bytes()
            r2 = self._run(d, _plan_rows("2024示例专业"),
                           {"outcome": "ok", "body": b"PK\x03\x04junk",
                            "size": 8,
                            "content_type": "application/msword"})
            after = paths["docx"].read_bytes()
        self.assertEqual(r2["status"], "partial")
        self.assertTrue(r2["items"])          # 返回旧成功条目，而非本次失败结果
        self.assertEqual(before, after)       # 假 PK 不得覆盖旧件

    def test_truncated_zip_rejected(self):
        tmp, d, _ = self._prep()
        body = _make_docx()
        with tmp:
            r = self._run(d, _plan_rows("2024示例专业"),
                          {"outcome": "ok", "body": body[:len(body) // 2],
                           "size": len(body) // 2,
                           "content_type": "application/msword"})
        self.assertEqual(r["status"], "unavailable")
        self.assertIn("校验失败", _message(r))

    def test_old_file_without_meta_not_reused(self):
        tmp, d, acct = self._prep()
        with tmp:
            svc = _svc(d)
            paths = svc._plan_doc_paths(acct, "2025", "2024示例专业")
            paths["dir"].mkdir(parents=True)
            paths["docx"].write_bytes(_make_docx())  # 只有同名文件，没有元数据
            r = self._run(d, _plan_rows("2024示例专业"),
                          {"outcome": "network_error", "diag": "synthetic"})
        self.assertEqual(r["status"], "unavailable")
        self.assertFalse(r.get("items"))

    def test_file_key_includes_account_and_grade(self):
        tmp, d, _ = self._prep()
        with tmp:
            svc = _svc(d)
            a = svc._plan_doc_paths("acct-a", "2025", "2024示例专业")
            b = svc._plan_doc_paths("acct-b", "2025", "2024示例专业")
            c = svc._plan_doc_paths("acct-a", "2024", "2024示例专业")
        self.assertNotEqual(a["docx"], b["docx"])
        self.assertNotEqual(a["docx"], c["docx"])
        self.assertIn("acct-a", str(a["docx"]))
        self.assertIn("2025", str(a["docx"]))

    def test_extract_docx_keeps_body_order(self):
        tmp, d, _ = self._prep()
        with tmp:
            p = d / "order.docx"
            p.write_bytes(_make_docx())
            text, complete = _extract_docx_text(p)
        lines = [x for x in text.split("\n") if x.strip()]
        i_a = lines.index("标题A")
        i_t = next(i for i, x in enumerate(lines) if "项目 | 值" in x)
        i_b = lines.index("标题B")
        i_t2 = next(i for i, x in enumerate(lines) if "条款 | 内容" in x)
        self.assertLess(i_a, i_t)
        self.assertLess(i_t, i_b)
        self.assertLess(i_b, i_t2)
        self.assertTrue(complete)

    def test_verify_rejects_html_with_pk_prefix(self):
        ok, why = _verify_docx_bytes(b"PK\x03\x04<html>not zip</html>")
        self.assertFalse(ok)
        self.assertTrue(why)


# ---------------- 5. 性能批：考试缓存 + 电表会话复用 ----------------

_METER_HTML = (
    "<script>const yffye = ref('35.30'); const bzdye = ref('0');"
    "const zye = ref('35.30'); const zydl = ref('10876.33');"
    "const dbzt = ref('保电'); const cbsj = ref('2026-10-18 02:44:34');"
    "const roomaddr = ref('1幢909'); const chargeType = ref('预付费');</script>")
_BILL_HTML = (
    "<script>billData = ref({qcdl: \"0\", qmdl: \"10\","
    " ydl: \"10\", ydje: \"5.5\"}); const month = ref('2026-10');</script>")


class _MeterSession:
    def __init__(self, fail_first_get: bool = False, fail_page: bool = False):
        self.posts = 0
        self.gets = 0
        self.closed = 0
        self.fail_first_get = fail_first_get
        self.fail_page = fail_page

    def post(self, url, **kwargs):
        self.posts += 1
        if self.fail_page and url.endswith("/electricdetail"):
            return _FakeResp(200, b"<html>errorApp</html>")
        return _FakeResp(200, b'{"url": "/meter_main"}')

    def get(self, url, **kwargs):
        self.gets += 1
        if self.fail_page:
            return _FakeResp(200, b"<html>errorApp</html>")
        if self.fail_first_get and self.gets == 1:
            raise E.requests.ConnectionError("synthetic reset")
        if url.endswith("/meter_main"):
            return _FakeResp(200, _METER_HTML.encode("utf-8"))
        if url.endswith("/electricbill"):
            return _FakeResp(200, _BILL_HTML.encode("utf-8"))
        raise AssertionError("unexpected URL: %s" % url)

    def close(self):
        self.closed += 1


class ElectricSessionReuseTests(unittest.TestCase):
    def setUp(self):
        E.clear_sessions()
        self.addCleanup(E.clear_sessions)

    def test_repeat_summary_bootstraps_and_selects_once(self):
        sess = _MeterSession()
        boot = mock.Mock(return_value=sess)
        with mock.patch.object(E, "_bootstrap", boot):
            r1 = E.fetch_meter_summary("M1")
            r2 = E.fetch_meter_summary("M1")
        self.assertEqual(r1["outcome"], "ok")
        self.assertEqual(r2["outcome"], "ok")
        self.assertEqual(boot.call_count, 1)   # 只建一次匿名会话
        self.assertEqual(sess.posts, 1)        # selectmeter 只发一次
        self.assertEqual(sess.gets, 2)         # 余额仍每次实取

    def test_bill_reuses_the_summary_session(self):
        sess = _MeterSession()
        boot = mock.Mock(return_value=sess)
        with mock.patch.object(E, "_bootstrap", boot):
            E.fetch_meter_summary("M1")
            r = E.fetch_meter_bill("M1")
        self.assertEqual(r["outcome"], "ok")
        self.assertEqual(boot.call_count, 1)
        self.assertEqual(sess.posts, 1)
        self.assertEqual(sess.gets, 2)

    def test_drop_forces_new_session_and_closes_old(self):
        first, second = _MeterSession(), _MeterSession()
        boot = mock.Mock(side_effect=[first, second])
        with mock.patch.object(E, "_bootstrap", boot):
            E.fetch_meter_summary("M1")
            E._drop_session("M1")
            r = E.fetch_meter_summary("M1")
        self.assertEqual(r["outcome"], "ok")
        self.assertEqual(boot.call_count, 2)
        self.assertEqual(first.closed, 1)

    def test_network_failure_drops_session_and_retries(self):
        bad = _MeterSession(fail_first_get=True)
        good = _MeterSession()
        boot = mock.Mock(side_effect=[bad, good])
        with mock.patch.object(E, "_bootstrap", boot):
            r = E.fetch_meter_summary("M1")
        self.assertEqual(r["outcome"], "ok")
        self.assertEqual(boot.call_count, 2)
        self.assertEqual(bad.closed, 1)


class ExamCacheWiringTests(unittest.TestCase):
    def test_get_exams_goes_through_query_cache(self):
        import src.mcp_server.server as server
        fake = {"status": "ok", "items": []}
        with mock.patch.object(server.query_cache, "call",
                               return_value=fake) as call:
            result = server.get_exams("2025", "3")
        call.assert_called_once()
        bound = call.call_args.args[0]
        self.assertIs(bound.__func__, server.jwweb_service.get_exams.__func__)
        self.assertEqual(call.call_args.kwargs["force_refresh"], False)
        self.assertEqual(result.structured_content, fake)

    def test_get_exams_exposes_force_refresh(self):
        from src.mcp_server import server
        tools = {t.name: t for t in asyncio.run(server.mcp.list_tools())}
        self.assertIn("force_refresh",
                      tools["get_exams"].input_schema["properties"])


# ---------------- 6. 存储事务与 WAL ----------------

class StoreTransactionTests(unittest.TestCase):
    def _store(self):
        tmp = temp_dir()
        self.addCleanup(tmp.cleanup)
        return tmp, SnapshotStore(Path(tmp.name) / "cache.sqlite3")

    def test_wal_and_synchronous_full(self):
        tmp, store = self._store()
        self.assertEqual(store.journal_mode(), "wal")
        con = store._conn()
        try:
            sync = con.execute("PRAGMA synchronous").fetchone()[0]
        finally:
            con.close()
        self.assertEqual(int(sync), 2)   # FULL，不做 NORMAL 的掉电回滚取舍

    def test_success_writes_snapshot_and_health_together(self):
        tmp, store = self._store()
        at = "2026-10-18T17:00:00+08:00"
        snap = store.upsert_snapshot_and_health(
            "acct-a", "m", "s", "scope", items=[{"x": 1}], fetched_at=at,
            complete=False, consistency="best_effort", success=True, at=at)
        self.assertTrue(snap["snapshot_id"])
        self.assertIsNotNone(store.get_snapshot("acct-a", "m", "s", "scope"))
        row = store.get_health_rows("acct-a")[0]
        self.assertEqual(row["last_success_at"], at)
        self.assertIsNone(row["last_issue"])

    def test_issue_does_not_advance_last_success(self):
        tmp, store = self._store()
        at = "2026-10-18T17:00:00+08:00"
        store.upsert_snapshot_and_health(
            "acct-a", "m", "s", "scope", items=[], fetched_at=at,
            complete=False, consistency="best_effort", success=False, at=at,
            issue={"code": "completeness_unproven"})
        row = store.get_health_rows("acct-a")[0]
        self.assertIsNone(row["last_success_at"])
        self.assertEqual(row["last_issue"]["code"], "completeness_unproven")

    def test_mid_write_failure_rolls_back_snapshot(self):
        tmp, store = self._store()
        with mock.patch.object(SnapshotStore, "_ensure_health_row",
                               side_effect=RuntimeError("synthetic")):
            with self.assertRaises(RuntimeError):
                store.upsert_snapshot_and_health(
                    "acct-a", "m", "s", "scope", items=[{"x": 1}],
                    fetched_at="2026-10-18T17:00:00+08:00", complete=False,
                    consistency="best_effort", success=True,
                    at="2026-10-18T17:00:00+08:00")
        self.assertIsNone(store.get_snapshot("acct-a", "m", "s", "scope"))
        self.assertEqual(store.get_health_rows("acct-a"), [])


# ---------------- 8. 批1复核新反例 ----------------

class GradeFieldTests(unittest.TestCase):
    def test_njdm_id_multi_mismatch_missing_conflict(self):
        from src.service.jwweb_service import _row_applies_to_grade
        self.assertTrue(_row_applies_to_grade({"njdm_id": "2024,2025"}, "2025")[0])
        self.assertFalse(_row_applies_to_grade({"njdm_id": "2024"}, "2025")[0])
        self.assertFalse(_row_applies_to_grade({}, "2025")[0])
        self.assertFalse(_row_applies_to_grade(
            {"njdm_id": "2024", "synj": "2025"}, "2025")[0])
        self.assertTrue(_row_applies_to_grade({"synj": "2025"}, "2025")[0])

    def test_non_year_token_rejected(self):
        from src.service.jwweb_service import _row_applies_to_grade
        ok, diag = _row_applies_to_grade({"njdm_id": "2025,garbage"}, "2025")
        self.assertFalse(ok)
        self.assertIn("token", diag)


class PlanPathTests(unittest.TestCase):
    def test_sanitized_names_do_not_collide(self):
        tmp, d, _ = TrainingPlanDocTests()._prep()
        with tmp:
            svc = _svc(d)
            a = svc._plan_doc_paths("acct-a", "2025", "SYNTHETIC/A")
            b = svc._plan_doc_paths("acct-a", "2025", "SYNTHETIC:A")
        self.assertNotEqual(a["docx"], b["docx"])


class PlanPublishTests(unittest.TestCase):
    def _prep(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        acct = _seed_jw(d)
        return tmp, d, acct

    def _downloads(self, bodies):
        counter = {"i": 0}
        lock = threading.Lock()

        def fake(transport, pid):
            with lock:
                body = bodies[counter["i"] % len(bodies)]
                counter["i"] += 1
            time.sleep(0.05)
            return {"outcome": "ok", "body": body, "size": len(body),
                    "content_type": "application/msword"}
        return fake

    def test_parallel_same_identity_keeps_one_consistent_version(self):
        tmp, d, acct = self._prep()
        bodies = [_make_docx("A"), _make_docx("B")]
        rows = _plan_rows("2024示例专业")
        with tmp:
            svc = _svc(d)
            with mock.patch.object(JWS, "fetch_training_plans",
                                   return_value=JwwebFetch("ok", rows, "", 1)):
                with mock.patch.object(JWS, "fetch_training_plan_doc",
                                       side_effect=self._downloads(bodies)):
                    with ThreadPoolExecutor(max_workers=2) as pool:
                        results = list(pool.map(
                            lambda _: svc.get_training_plan_doc("2024示例专业", "2025"),
                            range(2)))
            paths = svc._plan_doc_paths(acct, "2025", "2024示例专业")
            meta = json.loads(paths["meta"].read_text(encoding="utf-8"))
            actual = hashlib.sha256(paths["docx"].read_bytes()).hexdigest()
            leftovers = [p.name for p in paths["dir"].iterdir()
                         if ".tmp-" in p.name or ".bak-" in p.name]
        self.assertEqual([r["status"] for r in results], ["ok", "ok"])
        self.assertEqual(meta["sha256"], actual)
        self.assertEqual(leftovers, [])

    def test_rollback_failure_returns_unavailable_and_keeps_backup(self):
        tmp, d, acct = self._prep()
        first, second = _make_docx("A"), _make_docx("B")
        rows = _plan_rows("2024示例专业")
        with tmp:
            svc = _svc(d)
            with mock.patch.object(JWS, "fetch_training_plans",
                                   return_value=JwwebFetch("ok", rows, "", 1)):
                with mock.patch.object(JWS, "fetch_training_plan_doc",
                                       return_value={"outcome": "ok",
                                                     "body": first,
                                                     "size": len(first),
                                                     "content_type": "application/msword"}):
                    self.assertEqual(svc.get_training_plan_doc(
                        "2024示例专业", "2025")["status"], "ok")
                real_replace = os.replace
                calls = {"n": 0}

                def flaky(src, dst):
                    calls["n"] += 1
                    if calls["n"] in (2, 3):   # txt 替换失败 + 回滚 docx 失败
                        raise OSError("synthetic")
                    return real_replace(src, dst)

                with mock.patch.object(JWS, "fetch_training_plan_doc",
                                       return_value={"outcome": "ok",
                                                     "body": second,
                                                     "size": len(second),
                                                     "content_type": "application/msword"}):
                    with mock.patch.object(JWS.os, "replace", side_effect=flaky):
                        out = svc.get_training_plan_doc("2024示例专业", "2025")
            paths = svc._plan_doc_paths(acct, "2025", "2024示例专业")
            backups = [p for p in paths["dir"].iterdir() if ".bak-" in p.name]
        self.assertEqual(out["status"], "unavailable")
        self.assertIn("回滚未通过校验", _message(out))
        self.assertTrue(backups)


class SdtExtractionTests(unittest.TestCase):
    def test_content_control_is_extracted_in_order(self):
        from docx.oxml import OxmlElement
        tmp = temp_dir()
        with tmp:
            doc = Document()
            doc.add_paragraph("普通段落")
            p = doc.add_paragraph("控件条款")._p
            sdt = OxmlElement("w:sdt")
            content = OxmlElement("w:sdtContent")
            sdt.append(content)
            content.append(p)
            doc.element.body.append(sdt)
            path = Path(tmp.name) / "sdt.docx"
            doc.save(path)
            text, complete = _extract_docx_text(path)
        self.assertIn("控件条款", text)
        self.assertIn("普通段落", text)
        self.assertTrue(complete)


class ElectricRegressionTests(unittest.TestCase):
    def setUp(self):
        E.clear_sessions()
        self.addCleanup(E.clear_sessions)

    def _svc(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        (d / "dorm_meter.json").write_text(json.dumps({
            "meter": "M1", "building": "1幢", "floor": "9楼", "room": "909",
            "area_label": "示例学生公寓", "area_root_id": "001000099",
        }, ensure_ascii=False), encoding="utf-8")
        return tmp, d

    def test_bill_month_parsed_after_regex_rename(self):
        sess = _MeterSession()
        with mock.patch.object(E, "_bootstrap", return_value=sess):
            res = E.fetch_meter_bill("M1")
        self.assertEqual(res["outcome"], "ok")
        self.assertEqual(res["data"]["month"], "2026-10")

    def test_summary_selectmeter_array_is_controlled(self):
        class _ArraySession:
            def post(self, url, **kwargs):
                return _FakeResp(200, b"[]")

            def get(self, url, **kwargs):
                raise AssertionError("不得继续请求")

            def close(self):
                pass

        with mock.patch.object(E, "_bootstrap", return_value=_ArraySession()):
            res = E.fetch_meter_summary("M1", max_attempts=1)
        self.assertEqual(res["outcome"], "parse_error")

    def test_detail_default_month_uses_same_key(self):
        tmp, d = self._svc()
        month = ES.datetime.now(ES._TZ).strftime("%Y-%m")
        ok = {"outcome": "ok", "diag": "", "data": {
            "meter": "M1", "month": month, "days": [], "total_kwh": 0}}
        with tmp:
            svc = ElectricService(d)
            with mock.patch.object(ES.electric, "fetch_meter_detail", return_value=ok):
                svc.get_dorm_electric("detail")
                svc.get_dorm_electric("detail", month=month)
            entries = svc._load_all()
        self.assertEqual(len(entries), 1)   # 隐含本月与显式同月同键

    def test_location_change_invalidates_meter_binding(self):
        tmp, d = self._svc()
        with tmp:
            svc = ElectricService(d)
            svc._meter(svc._load_config())          # 先建立绑定
            cfg = svc._load_config()
            cfg["room"] = "404"
            meter, err = svc._meter(cfg)
        self.assertIsNone(meter)
        self.assertIn("不一致", err)

    def test_legacy_config_migrates_location_binding(self):
        tmp, d = self._svc()
        with tmp:
            svc = ElectricService(d)
            meter, err = svc._meter(svc._load_config())
            saved = svc._load_config()
        self.assertEqual(meter, "M1")
        self.assertIsNone(err)
        self.assertTrue(saved.get("location_key"))


class SeatAccountContextTests(unittest.TestCase):
    def test_cookie_uses_the_same_account_snapshot(self):
        tmp = temp_dir()
        d = Path(tmp.name)
        seed_session(d, "acct-a")
        sess = S.load_session(d, "acct-a") or {}
        sess["seat_cookie"] = "wechatSESS_ID=A"
        S.save_session(d, "acct-a", sess)
        payload = {"areas": [], "by_floor": {}, "rank_by_rate": []}
        with tmp:
            with mock.patch.object(SS.S, "get_active",
                                   side_effect=["acct-a", "acct-b"]) as active:
                with mock.patch.object(SS, "fetch_seat_summary",
                                       return_value={"outcome": "ok", "diag": "",
                                                     "data": payload}):
                    r = SeatService(d).get_seat_summary()
            state = json.loads((d / "state" / "seat_last.json").read_text("utf-8"))
        self.assertEqual(r["status"], "ok")
        self.assertEqual(active.call_count, 1)
        self.assertEqual(state["account"], "acct-a")


class ElectricDeadSessionTests(unittest.TestCase):
    """§20.4-1：bill/detail 遇失效页必须淘汰池中会话并有限重建。"""

    def setUp(self):
        E.clear_sessions()
        self.addCleanup(E.clear_sessions)

    def _factory(self, made):
        def bootstrap(timeout=20.0):
            s = _MeterSession()
            s.fail_page = True
            made.append(s)
            return s
        return bootstrap

    def test_bill_error_page_drops_session_and_rebuilds(self):
        made = []
        with mock.patch.object(E, "_bootstrap", side_effect=self._factory(made)):
            r1 = E.fetch_meter_bill("M1")
            self.assertEqual(r1["outcome"], "error")
            self.assertEqual(len(made), 2)         # 失效页 → 淘汰 + 重建一次
            r2 = E.fetch_meter_bill("M1")
        self.assertEqual(r2["outcome"], "error")
        self.assertEqual(len(made), 4)             # 下一查询同样有限重建
        self.assertTrue(all(s.closed == 1 for s in made))

    def test_detail_error_page_drops_session_and_rebuilds(self):
        made = []
        with mock.patch.object(E, "_bootstrap", side_effect=self._factory(made)):
            r = E.fetch_meter_detail("M1")
        self.assertEqual(r["outcome"], "error")
        self.assertEqual(len(made), 2)             # 失效页 → 淘汰 + 重建一次
        self.assertTrue(all(s.closed == 1 for s in made))


class ElectricMigrationEvidenceTests(unittest.TestCase):
    def _svc_with_state(self, room="909", building="1幢"):
        tmp = temp_dir()
        d = Path(tmp.name)
        (d / "dorm_meter.json").write_text(json.dumps({
            "meter": "M1", "building": building, "floor": "9楼", "room": room,
            "area_label": "示例学生公寓", "area_root_id": "001000099",
        }, ensure_ascii=False), encoding="utf-8")
        return tmp, d

    def test_migration_refuses_state_contradiction(self):
        tmp, d = self._svc_with_state()
        with tmp:
            svc = ElectricService(d)
            svc._save_last(svc._key("summary", None, "M1"), "M1",
                           {"data": {"room": {"room": "909"}}},
                           "2026-10-18T12:00:00+08:00")
            cfg = svc._load_config()
            cfg["room"] = "404"
            cfg.pop("location_key", None)
            meter, err = svc._meter(cfg)
        self.assertIsNone(meter)
        self.assertIn("矛盾", err)

    def test_migration_refuses_same_room_other_building(self):
        # §22 唯一 P2：同房号换楼栋/楼层/园区也必须被拦下
        tmp, d = self._svc_with_state(room="909")
        with tmp:
            svc = ElectricService(d)
            svc._save_last(svc._key("summary", None, "M1"), "M1",
                           {"data": {"room": {"building": "1幢", "floor": "9楼",
                                              "room": "909",
                                              "area_root_id": "001000099"}}},
                           "2026-10-18T12:00:00+08:00")
            cfg = svc._load_config()
            cfg["building"] = "31幢"           # 房号相同，楼栋不同
            cfg.pop("location_key", None)
            meter, err = svc._meter(cfg)
        self.assertIsNone(meter)
        self.assertIn("building", err)

    def test_migration_refuses_same_room_other_campus(self):
        tmp, d = self._svc_with_state(room="909")
        with tmp:
            svc = ElectricService(d)
            svc._save_last(svc._key("summary", None, "M1"), "M1",
                           {"data": {"room": {"building": "1幢", "floor": "9楼",
                                              "room": "909",
                                              "area_root_id": "001000099"}}},
                           "2026-10-18T12:00:00+08:00")
            cfg = svc._load_config()
            cfg["area_root_id"] = "001000010"  # 房号相同，园区不同
            cfg.pop("location_key", None)
            meter, err = svc._meter(cfg)
        self.assertIsNone(meter)
        self.assertIn("area_root_id", err)

    def test_migration_passes_when_fields_agree(self):
        tmp, d = self._svc_with_state(room="909")
        with tmp:
            svc = ElectricService(d)
            svc._save_last(svc._key("summary", None, "M1"), "M1",
                           {"data": {"room": {"building": "1幢", "floor": "9楼",
                                              "room": "909",
                                              "area_root_id": "001000099"}}},
                           "2026-10-18T12:00:00+08:00")
            cfg = svc._load_config()
            cfg.pop("location_key", None)
            meter, err = svc._meter(cfg)
            saved = svc._load_config()
        self.assertEqual(meter, "M1")
        self.assertIsNone(err)
        self.assertTrue(saved.get("location_key"))


class MaintenanceGuardTests(unittest.TestCase):
    def test_backup_refuses_same_file_and_existing_dest(self):
        import sqlite3
        from src.storage import maintenance as M
        tmp = temp_dir()
        self.addCleanup(tmp.cleanup)
        d = Path(tmp.name)
        store = SnapshotStore(d / "cache.sqlite3")
        store.upsert_snapshot("a", "m", "s", "scope", [{"x": 1}],
                              "2026-10-18T17:00:00+08:00", False, "best_effort")
        db = d / "cache.sqlite3"
        with self.assertRaises(ValueError):
            M.backup(db, db)                       # 拒绝自备份
        dest = d / "backup.sqlite3"
        dest.write_bytes(b"not-a-db")
        with self.assertRaises(FileExistsError):
            M.backup(db, dest)                     # 默认拒绝覆盖
        info = M.backup(db, dest, overwrite=True)  # 显式 --force 才覆盖
        self.assertGreater(info["size_bytes"], 0)
        con = sqlite3.connect(str(dest))
        try:
            self.assertEqual(con.execute(
                "SELECT COUNT(*) FROM snapshots").fetchone()[0], 1)
        finally:
            con.close()


# ---------------- 10. 缓存维护脚本 ----------------

class MaintenanceTests(unittest.TestCase):
    def test_backup_checkpoint_status(self):
        import sqlite3
        from src.storage import maintenance as M
        tmp = temp_dir()
        self.addCleanup(tmp.cleanup)
        d = Path(tmp.name)
        store = SnapshotStore(d / "cache.sqlite3")
        store.upsert_snapshot("acct-a", "m", "s", "scope", [{"x": 1}],
                              "2026-10-18T17:00:00+08:00", False, "best_effort")
        info = M.backup(d / "cache.sqlite3", d / "backup.sqlite3")
        self.assertGreater(info["size_bytes"], 0)
        con = sqlite3.connect(str(d / "backup.sqlite3"))
        try:
            self.assertEqual(con.execute(
                "SELECT COUNT(*) FROM snapshots").fetchone()[0], 1)
        finally:
            con.close()
        ck = M.checkpoint(d / "cache.sqlite3")
        self.assertFalse(ck["busy"])
        self.assertEqual(ck["wal_bytes_after"], 0)
        st = M.status(d / "cache.sqlite3")
        self.assertEqual(st["journal_mode"], "wal")
        self.assertEqual(st["snapshots"], 1)
        self.assertEqual(st["source_health"], 0)


# ---------------- 11. 输出 schema ----------------

class PreserveCredentialsTests(unittest.TestCase):
    """重跑 init 不得清掉独立凭据（选座 Cookie / 体测令牌）。"""

    def test_merge_keeps_independent_credentials_only(self):
        from src.auth.auto_login import _preserve_independent_credentials
        new = {"zjgsu_token": "new"}
        old = {"zjgsu_token": "old", "jwglxt_cookies": [{"name": "x"}],
               "seat_cookie": "wechatSESS_ID=SECRET",
               "seat_cookie_saved_at": "2026-10-18T12:07:14+08:00",
               "tzjk_token": "T", "tzjk_student_no": "0000000000"}
        kept = _preserve_independent_credentials(new, old)
        self.assertIn("seat_cookie", kept)
        self.assertIn("tzjk_token", kept)
        self.assertNotIn("jwglxt_cookies", new)   # 统一认证会话必须换新
        self.assertEqual(new["zjgsu_token"], "new")

    def test_absent_keys_are_not_invented(self):
        from src.auth.auto_login import _preserve_independent_credentials
        new = {"zjgsu_token": "new"}
        self.assertEqual(_preserve_independent_credentials(new, None), [])
        self.assertNotIn("seat_cookie", new)


class OutputSchemaTests(unittest.TestCase):
    def test_two_new_tools_have_output_schema(self):
        from src.mcp_server import server
        tools = {t.name: t for t in asyncio.run(server.mcp.list_tools())}
        for name in ("get_library_seats", "get_dorm_electric"):
            self.assertTrue(getattr(tools[name], "output_schema", None),
                            "%s 缺少 outputSchema" % name)


if __name__ == "__main__":
    unittest.main()
